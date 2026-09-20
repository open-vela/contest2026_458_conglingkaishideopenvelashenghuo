#!/usr/bin/env python3
"""对技术报告 .docx 做事实性勘误（不改版式）。

改错不改风格：只在段落文本层做精确替换，命中跨 run 时把整段文本收敛到
第一个 run（报告正文没有混排字号，视觉无差别）。

    python fix_report.py <in.docx> <out.docx>
"""
import sys

import docx

# (旧, 新, 说明)
FIXES = [
    ("0x09(REMINDER)", "0x12(REMINDER)",
     "帧号写错：0x09 是 ACK，提醒主动推送是 0x12"),
    ("本地执行或 0x02 上行云端 → 0x03 行协议",
     "本地执行或 0x04(CLOUD_REQ) 上行云端 → 0x05(CLOUD_RESP) 行协议",
     "云端上下行帧号写错（0x02/0x03 是预留的音频帧）"),
    ("网关端为 Python 包 wristclaw-host——",
     "网关端为 Python 包 host/——",
     "目录已改名为 host/"),
    ("bridge.py 读取 wristclaw-host/skills/*.md",
     "bridge.py 读取 host/skills/*.md",
     "同上"),
    ("app/wristclaw 板端 C 代码 + wristclaw-host 上位机代码",
     "app/wristclaw 板端 C 代码 + host 上位机代码",
     "同上"),
]


def replace_in_paragraph(p, old, new):
    if old not in p.text:
        return False
    # 先试单 run 精确替换（保留 run 级格式）
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return True
    # 跨 run：收敛到第一个 run
    full = p.text.replace(old, new)
    for i, run in enumerate(p.runs):
        run.text = full if i == 0 else ''
    return True


def main():
    src, dst = sys.argv[1], sys.argv[2]
    doc = docx.Document(src)

    hit = {old: 0 for old, _, _ in FIXES}

    def sweep(paragraphs):
        for p in paragraphs:
            for old, new, _ in FIXES:
                if replace_in_paragraph(p, old, new):
                    hit[old] += 1
                    print("  [fix] %s  ->  %s" % (old[:46], new[:46]))

    sweep(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                sweep(cell.paragraphs)
                for tt in cell.tables:
                    for r2 in tt.rows:
                        for c2 in r2.cells:
                            sweep(c2.paragraphs)

    print("\n替换统计：")
    missing = []
    for old, _, why in FIXES:
        print("  %-4s %s   (%s)" % (hit[old] or "MISS", why, old[:40]))
        if not hit[old]:
            missing.append(old)

    doc.save(dst)
    print("\nwrote %s" % dst)
    return 1 if missing else 0


if __name__ == '__main__':
    sys.exit(main())
